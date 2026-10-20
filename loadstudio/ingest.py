"""Bounded deterministic adapters; unsupported input stays visible for review."""
from __future__ import annotations
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
import hashlib
import html
import io
import re
import zipfile
from .models import Load, Place, Window
from .engine import COUNTRIES, normalized

MONTHS={m:i for i,vs in enumerate([('jan','ene'),('feb',),('mar',),('apr','abr'),('may',),('jun',),('jul',),('aug','ago'),('sep',),('oct',),('nov',),('dec','dic')],1) for m in vs}
GOODS={'OIL':('Aceite','Oil'),'WATER':('Agua','Water'),'ACEITE':('Aceite','Oil'),'AGUA':('Agua','Water'),
       'ROPA':('Ropa','Clothing'),'CLOTHING':('Ropa','Clothing')}

def numeric(text):
    s=str(text).strip().replace('\u00a0','').replace(' ','')
    if ',' in s:
        if s.count(',')!=1: raise ValueError('Separadores numericos ambiguos')
        s=s.replace('.','').replace(',','.')
    if not re.fullmatch(r'[0-9]+(?:\.[0-9]+)?',s):raise ValueError('Numero no reconocido')
    value=Decimal(s)
    if not value.is_finite():raise ValueError('Numero no finito')
    return value

def parse_window(text):
    s=text.strip(); d=None
    iso=re.search(r'\b([0-9]{4})-([0-9]{2})-([0-9]{2})\b',s)
    eu=re.search(r'\b([0-9]{1,2})/([0-9]{1,2})/([0-9]{4})\b',s)
    word=re.search(r'\b([0-9]{1,2})\s+([A-Za-z]{3,9}),?\s+([0-9]{4})\b',s)
    span=None
    if iso:d=date(*map(int,iso.groups()));span=iso.span()
    elif eu:d=date(int(eu[3]),int(eu[2]),int(eu[1]));span=eu.span()
    elif word and word[2][:3].lower() in MONTHS:d=date(int(word[3]),MONTHS[word[2][:3].lower()],int(word[1]));span=word.span()
    if d is None:raise ValueError('Fecha explicita con ano necesaria; no se inventa el ano')
    tail=s[span[1]:]; times=re.findall(r'\b([0-9]{1,2}):([0-9]{2})\b',tail)
    if len(times)>2:raise ValueError('Mas de dos horas en una ventana')
    if not times:
        shorthand=re.search(r'\b([0-9]{1,2})\s*[-\u2013]\s*([0-9]{1,2})\b',tail)
        if shorthand:times=[(shorthand[1],'00'),(shorthand[2],'00')]
    ts=[time(int(h),int(m)) for h,m in times]
    return Window(date=d,start=ts[0] if ts else None,end=ts[1] if len(ts)>1 else None)

def parse_place(text):
    s=' '.join(text.split());p=Place()
    iso=re.search(r'\b(ES|PT|FR|DE|IT|NL|BE|GB)\b',s,re.I)
    if iso:p.country=COUNTRIES[iso[1].upper()];s=s[:iso.start()]+s[iso.end():]
    cp=re.search(r'\b([0-9]{4}-[0-9]{3}|[0-9]{5}|[0-9]{4})\b',s)
    if cp:p.postcode=cp[1];s=s[:cp.start()]+s[cp.end():]
    s=re.sub(r'[,()]',' ',s).strip();s=' '.join(s.split())
    if s:p.locality=s
    return p

def clean(text):
    return '\n'.join(x.strip(' \\') for x in html.unescape(text).replace('\r','').split('\n'))

def fote(text):
    s=clean(text);lines=[x.strip() for x in s.splitlines() if x.strip()];data={};issues=[]
    def warning(field,message):issues.append({'field':field,'message':message,'severity':'warning'})
    def after(label):
        values=[]
        for i,line in enumerate(lines):
            if normalized(line).rstrip(':')==normalized(label).rstrip(':') and i+1<len(lines):values.append(lines[i+1])
            elif line.lower().startswith(label.lower()+':') and line[len(label)+1:].strip():values.append(line[len(label)+1:].strip())
        return values
    def one(field,values):
        values=list(dict.fromkeys(values))
        if len(values)>1:warning(field,'Valores contradictorios en FOTE: completar el campo manualmente');return None
        return values[0] if values else None
    refs=re.findall(r'Order\s*#\s*:?\s*([0-9]+)',s,re.I)
    data['reference']=one('reference',refs)
    for field,label in [('pickup','Pick Up'),('delivery','Delivery')]:
        candidates=[x for x in after(label) if re.search(r'[0-9]{4}',x)]
        val=one(field,candidates)
        if val:
            try:data[field]=parse_window(val)
            except ValueError as e:warning(field,str(e))
    for field,start,end in [('origin','Origin','Destination'),('destination','Destination','Related Orders')]:
        try:
            a=next(i for i,l in enumerate(lines) if normalized(l)==normalized(start))
            b=next((i for i in range(a+1,len(lines)) if normalized(lines[i]) in {normalized(end),'CUSTOMER','CARRIER'}),len(lines))
            block=lines[a+1:b];p=Place()
            for i,line in enumerate(block):
                if line.upper() in COUNTRIES:p.country=COUNTRIES[line.upper()]
                match=re.fullmatch(r'(.+?),\s*[0-9]{1,3}\s+([0-9]{4}-[0-9]{3}|[0-9]{5}|[0-9]{4})',line)
                if match:p.locality=match[1].strip();p.postcode=match[2]
                if normalized(line).rstrip(':')=='LOCATION' and i+1<len(block):p.center=block[i+1]
            data[field]=p
        except StopIteration:warning(field,'Bloque de direccion no localizado')
    customers=[]
    for i,line in enumerate(lines):
        if normalized(line)=='CUSTOMER':
            block=lines[i+1:i+5]
            for j,l in enumerate(block):
                if normalized(l).rstrip(':')=='NAME' and j+1<len(block):customers.append(block[j+1])
    data['customer']=one('customer',customers)
    data['vehicle']=one('vehicle',after('Equipment Type'))
    service=one('service_code',after('Customer Service Type'))
    if service:
        match=re.match(r'(T100|TR014|NTIE1|NT214|NT014)\b',service)
        if match:data['service_code']=match[1]
    weights=re.findall(r'([0-9][0-9 .,]*)\s*kg\(s\)',s,re.I)
    try:
        value=one('weight_kg',[str(numeric(w)) for w in weights])
        if value is not None:data['weight_kg']=value
    except ValueError as e:warning('weight_kg',str(e))
    revenue=one('client_price_eur',after('Revenue'))
    if revenue:
        try:data['client_price_eur']=numeric(revenue.replace('\u20ac','').replace('EUR',''))
        except ValueError as e:warning('client_price_eur',str(e))
    goods=one('goods_es',after('Commodity'))
    if goods:
        if normalized(goods) in GOODS:data['goods_es'],data['goods_en']=GOODS[normalized(goods)]
        else:data['goods_es']=goods;warning('goods_en','Mercancia original conservada; falta traduccion verificada')
    meters=one('linear_meters',after('Linear Meters'))
    if meters:
        # FOTE displays a row of original/final/actual values; all zero is its empty placeholder.
        values=re.findall(r'[0-9]+(?:[.,][0-9]+)?',meters)
        values=[numeric(v) for v in values]
        unique=set(values)
        if len(unique)==1 and values[0]>0:data['linear_meters']=values[0]
        elif len(unique)>1:warning('linear_meters','Metros lineales contradictorios; revisar original/final/actual')
    stop_match=re.search(r'\bStops\s*([0-9]+)\b',s,re.I)
    if stop_match and int(stop_match[1])>2:
        data['stop_count']=int(stop_match[1]);warning('stop_count','Hay mas de dos paradas: no exportar como ruta directa')
    return data,issues

LABELS={'REFERENCIA':'reference','REF':'reference','CLIENTE':'customer','CUSTOMER':'customer','ORIGEN':'origin','ORIGIN':'origin',
        'DESTINO':'destination','DESTINATION':'destination','RECOGIDA':'pickup','PICKUP':'pickup','CARGA':'pickup',
        'DESCARGA':'delivery','DELIVERY':'delivery','PESO':'weight_kg','WEIGHT':'weight_kg','ML':'linear_meters',
        'METROS':'linear_meters','VEHICULO':'vehicle','VEHICLE':'vehicle','SERVICIO':'service','SERVICE':'service',
        'PRECIO CLIENTE':'client_price_eur','CLIENT PRICE':'client_price_eur','MERCANCIA':'goods_es','GOODS':'goods_es','ADR':'adr',
        'PALETIZADO':'palletized','PALETS':'pallets','PARADAS':'stop_count','STOPS':'stop_count'}

def free_text(text):
    s=clean(text);data={};issues=[];seen={}
    def warn(f,m):issues.append({'field':f,'message':m,'severity':'warning'})
    route=re.search(r'(?m)^([^\n:]+?)\s*(?:->|\u2192)\s*([^\n]+)$',s)
    if route:
        route_parts=re.split(r'(?:->|\u2192)',route[0])
        data['origin']=parse_place(route_parts[0]);data['destination']=parse_place(route_parts[-1])
        if len(route_parts)>2:
            data['stop_count']=len(route_parts);warn('stop_count','Ruta con paradas intermedias: completar acciones, lugares y fechas antes de copiar')
    for line in re.split(r'[\n;]+',s):
        if not line.strip():continue
        pair=re.match(r'^\s*([^:]+):\s*(.*)$',line)
        if not pair:
            if not ('->' in line or '\u2192' in line):warn('input','Texto no interpretado: '+line[:100])
            continue
        field=LABELS.get(normalized(pair[1]));value=pair[2].strip()
        if not field:warn('input','Etiqueta no soportada: '+pair[1]);continue
        if field in seen and seen[field]!=value:data.pop(field,None);warn(field,'Dos valores distintos: revisar');seen[field]='CONFLICT';continue
        seen[field]=value
        try:
            if field in ('origin','destination'):data[field]=parse_place(value)
            elif field in ('pickup','delivery'):data[field]=parse_window(value)
            elif field=='weight_kg':
                unit=re.fullmatch(r'([0-9][0-9 .,]*)\s*(kg|t|toneladas?)',value,re.I)
                if not unit:raise ValueError('Peso necesita unidad kg o t')
                data[field]=numeric(unit[1])*(1 if unit[2].lower()=='kg' else 1000)
            elif field in ('linear_meters','client_price_eur'):data[field]=numeric(re.sub(r'(\s*m|\s*\u20ac|\s*EUR)$','',value,flags=re.I))
            elif field in ('adr','palletized'):
                lookup={'S':True,'SI':True,'YES':True,'N':False,'NO':False}
                if normalized(value) not in lookup:raise ValueError('Usar S o N; vacio no significa N')
                data[field]=lookup[normalized(value)]
            elif field=='service':
                options={'FTL':'FTL','GRUPAJE':'GROUPAGE','GROUPAGE':'GROUPAGE','RIGID':'RIGID'}
                if normalized(value) not in options:raise ValueError('Servicio no reconocido')
                data[field]=options[normalized(value)]
            elif field in ('pallets','stop_count'):data[field]=int(value)
            elif field=='goods_es':
                if '/' in value:data['goods_es'],data['goods_en']=[x.strip() for x in value.split('/',1)]
                elif normalized(value) in GOODS:data['goods_es'],data['goods_en']=GOODS[normalized(value)]
                else:data[field]=value
            else:data[field]=value
        except (ValueError,InvalidOperation) as e:warn(field,str(e))
    return data,issues

def parse(text,mode='auto'):
    use_fote=mode=='fote' or (mode=='auto' and re.search(r'Order\s*#',text,re.I))
    # Multiple complete copied orders are separate drafts, never merged by route.
    parts=re.split(r'(?im)(?=^\s*Order\s*#\s+[0-9]+\s+)',clean(text)) if use_fote else [text]
    parts=[p for p in parts if re.search(r'Order\s*#',p,re.I)] if use_fote else parts
    if not parts:parts=[text]
    if len(parts)>50:raise ValueError('Maximo 50 cargas por lote')
    out=[]
    for part in parts:
        data,issues=(fote(part) if use_fote else free_text(part))
        data['blocked_fields']=list(dict.fromkeys(i['field'] for i in issues if i['field'] not in ('input','stops','stop_count')))
        data['source_hash']=hashlib.sha256(part.encode()).hexdigest()
        try:load=Load.model_validate(data)
        except ValueError as e:
            # Retain all extractable fields. Invalid windows have already been isolated by adapters.
            issues.append({'field':'input','message':'Conflicto de datos: '+str(e)[:240],'severity':'error'})
            if 'pickup' in data and 'delivery' in data:
                data.pop('delivery');data['blocked_fields'].append('delivery.date');load=Load.model_validate(data)
            else:raise
        out.append({'load':load.model_dump(mode='json'),'issues':issues,'adapter':'fote' if use_fote else 'text'})
    return out

FIELD_MAP={'Referencia':'reference','Fecha disponibilidad':'pickup.date','Pa\u00eds origen':'origin.country','CP origen':'origin.postcode',
 'Localidad origen':'origin.locality','Pa\u00eds destino':'destination.country','CP destino':'destination.postcode','Localidad destino':'destination.locality',
 'Tipo Bolsa':'service','Tipo Vehiculo':'vehicle','Paradas':'stop_count','Stops':'stop_count','Peso':'weight_t','Largo':'linear_meters','Adr':'adr','Paletizado':'palletized'}
IMPORT_FIELDS={'reference','customer','origin.country','origin.postcode','origin.locality','origin.center','destination.country','destination.postcode','destination.locality',
 'pickup.date','pickup.start','pickup.end','delivery.date','delivery.start','delivery.end','service','vehicle','weight_kg','weight_t','linear_meters','adr','palletized','goods_es','goods_en','client_price_eur','stop_count'}

def workbook_preview(raw):
    if len(raw)>5_000_000:raise ValueError('Excel demasiado grande (maximo 5 MB)')
    if not raw.startswith(b'PK'):raise ValueError('Importacion general: usar .xlsx. El auditor BIFF8 es solo para verificar la plantilla, no un importador XLS universal.')
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist())>30_000_000 or len(z.infolist())>1000:raise ValueError('Excel descomprimido demasiado grande')
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    result=[]
    try:
        for sheet in book.worksheets[:10]:
            rows=[]
            for i,row in enumerate(sheet.iter_rows(max_row=52,max_col=40,values_only=True)):
                vals=[v.isoformat() if isinstance(v,(date,datetime,time)) else v for v in row]
                while vals and vals[-1] is None:vals.pop()
                if vals:rows.append(vals)
            result.append({'sheet':sheet.title,'rows':rows,'limit':50,'truncated':sheet.max_row>51})
    finally:book.close()
    return result

def import_rows(rows,mapping):
    if len(rows)>50:raise ValueError('Maximo 50 filas')
    if set(mapping)-IMPORT_FIELDS:raise ValueError('Mapeo con campos desconocidos')
    if len(set(mapping.values()))!=len(mapping):raise ValueError('Una columna asignada a dos campos')
    if any(not isinstance(i,int) or i<0 or i>=40 for i in mapping.values()):raise ValueError('Indice de columna invalido')
    out=[]
    for row in rows:
        if not any(v not in ('',None) for v in row):continue
        data={};issues=[]
        for field,index in mapping.items():
            value=row[index] if index<len(row) else None
            if value is None or value=='':continue
            if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):
                issues.append({'field':field,'message':'Formula rechazada: copiar un valor verificado','severity':'error'});continue
            if field.endswith('postcode') and not isinstance(value,str):
                issues.append({'field':field,'message':'CP numerico: posibles ceros perdidos; revisar','severity':'error'});continue
            try:
                if field in ('weight_kg','weight_t','linear_meters','client_price_eur'):
                    value=numeric(value)
                    if field=='weight_t':field='weight_kg';value*=1000
                elif field=='stop_count':
                    value=numeric(value)
                    if value!=value.to_integral_value():raise ValueError('Numero de paradas debe ser entero')
                    value=int(value)
                elif field in ('adr','palletized'):
                    if str(value).upper() not in ('S','N'):raise ValueError('S/N invalido')
                    value=str(value).upper()=='S'
                elif field=='service':value={'Tr\u00e1ilers Completos':'FTL','Grupajes':'GROUPAGE','R\u00edgidos Completos':'RIGID'}.get(value,value)
                elif field.endswith('.date'):
                    try:value=date.fromisoformat(str(value)[:10]).isoformat()
                    except ValueError:value=parse_window(str(value)).date.isoformat()
                else:value=str(value)
                if '.' in field:
                    group,key=field.split('.',1);data.setdefault(group,{})[key]=value
                else:data[field]=value
            except ValueError as e:issues.append({'field':field,'message':str(e),'severity':'error'})
        data['blocked_fields']=[i['field'] for i in issues]
        try:out.append({'load':Load.model_validate(data).model_dump(mode='json'),'issues':issues,'adapter':'xlsx'})
        except ValueError as e:out.append({'load':Load().model_dump(mode='json'),'issues':[{'field':'row','message':str(e)[:300],'severity':'error'}],'adapter':'xlsx'})
    return out
